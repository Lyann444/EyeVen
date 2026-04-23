import csv
import io
import logging
import os

import openpyxl

logger = logging.getLogger(__name__)


class ParserService:
    @staticmethod
    def parse_csv(file_path):
        if not os.path.isfile(file_path):
            return {"error": "File not found", "rows": [], "headers": [], "row_count": 0}

        try:
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                headers = reader.fieldnames or []
                rows = [row for row in reader]

            return {
                "headers": headers,
                "rows": rows,
                "row_count": len(rows),
                "status": "completed",
            }
        except Exception as e:
            logger.error("CSV parse error for %s: %s", file_path, e)
            return {"error": str(e), "rows": [], "headers": [], "row_count": 0, "status": "failed"}

    @staticmethod
    def parse_xlsx(file_path):
        if not os.path.isfile(file_path):
            return {"error": "File not found", "sheets": {}, "status": "failed"}

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheets = {}

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                headers = []

                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i == 0:
                        headers = [str(c) if c is not None else f"column_{j}" for j, c in enumerate(row)]
                    else:
                        row_dict = {}
                        for j, val in enumerate(row):
                            key = headers[j] if j < len(headers) else f"column_{j}"
                            row_dict[key] = str(val) if val is not None else ""
                        rows.append(row_dict)

                sheets[sheet_name] = {
                    "headers": headers,
                    "rows": rows,
                    "row_count": len(rows),
                }

            wb.close()

            return {
                "sheets": sheets,
                "sheet_names": list(sheets.keys()),
                "status": "completed",
            }
        except Exception as e:
            logger.error("XLSX parse error for %s: %s", file_path, e)
            return {"error": str(e), "sheets": {}, "status": "failed"}

    @staticmethod
    def parse_file(file_path):
        ext = os.path.splitext(file_path)[1].lstrip(".").lower()
        if ext == "csv":
            return ParserService.parse_csv(file_path)
        elif ext in ("xlsx", "xls"):
            return ParserService.parse_xlsx(file_path)
        else:
            return {"error": f"Unsupported structured format: {ext}", "status": "skipped"}
